from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "局部放电声学成像系统技术文档.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
RED = "9B1C1C"
GOLD = "7A5A00"
CONTENT_WIDTH = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width))
    tcW.set(qn("w:type"), "dxa")
    cell.width = Inches(width / 1440.0)


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for gridcol, width in zip(grid.gridCol_lst, widths):
        gridcol.set(qn("w:w"), str(width))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margin(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_paragraph_border(paragraph, bottom_color=BLUE, size="12", space="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), bottom_color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_keep(paragraph, keep_with_next=False, keep_together=False):
    pPr = paragraph._p.get_or_add_pPr()
    if keep_with_next:
        pPr.append(OxmlElement("w:keepNext"))
    if keep_together:
        pPr.append(OxmlElement("w:keepLines"))


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, size=10.5, color=INK, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, size=10.5, color="222222")
    else:
        r = p.add_run(text)
        set_font(r, size=10.5, color="222222")
    set_keep(p, keep_together=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_font(r, size=10.5, color="222222")
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_font(r, size=10.5, color="222222")
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 11.5}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    p.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 8}[level])
    p.paragraph_format.space_after = Pt({1: 8, 2: 6, 3: 4}[level])
    set_keep(p, keep_with_next=True)
    return p


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH - 240])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF2F6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for line_no, line in enumerate(text.splitlines()):
        if line_no:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, name="Consolas", size=8.6, color="263238")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_note(doc, title, text, kind="note"):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH - 240])
    cell = table.cell(0, 0)
    color = CALLOUT if kind == "note" else ("FFF7E6" if kind == "warning" else "FDEEEE")
    accent = DARK_BLUE if kind == "note" else (GOLD if kind == "warning" else RED)
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    label = p.add_run(title + "  ")
    set_font(label, size=10.5, color=accent, bold=True)
    body = p.add_run(text)
    set_font(body, size=10.5, color="222222")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header_cells = table.rows[0].cells
    for i, text in enumerate(headers):
        set_cell_shading(header_cells[i], LIGHT_GRAY)
        p = header_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=9.5, color=INK, bold=True)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_font(r, size=9.2, color="222222")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_flow(doc, items):
    for index, (label, body) in enumerate(items):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.05)
        r = p.add_run(label)
        set_font(r, size=10.5, color=DARK_BLUE, bold=True)
        r = p.add_run("  " + body)
        set_font(r, size=10.5, color="222222")
        if index != len(items) - 1:
            arrow = doc.add_paragraph()
            arrow.paragraph_format.left_indent = Inches(0.18)
            arrow.paragraph_format.space_after = Pt(3)
            r = arrow.add_run("|")
            set_font(r, name="Consolas", size=9, color=BLUE, bold=True)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def set_page_number(paragraph):
    run = paragraph.add_run("第 ")
    set_font(run, size=8.5, color="666666")
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run = paragraph.add_run(" 页")
    set_font(run, size=8.5, color="666666")


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for level, size, color in ((1, 16, BLUE), (2, 13, BLUE), (3, 11.5, DARK_BLUE)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hr = hp.add_run("局部放电声学成像系统技术文档")
    set_font(hr, size=8.5, color="666666")
    set_paragraph_border(hp, "D6DEE8", "4", "4")
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_page_number(fp)


def build_document():
    doc = Document()
    configure_document(doc)

    # Title block
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("技术设计与实施说明")
    set_font(r, size=11, color=DARK_BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("局部放电声学成像系统")
    set_font(r, size=24, color=INK, bold=True)
    set_paragraph_border(p, BLUE, "14", "10")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("64 通道声学阵列 | Zynq-7020 | FPGA 频域定位 | UDP 云图与原始时域波形")
    set_font(r, size=12, color="555555")

    add_table(doc, ["项目", "内容"], [
        ("文档版本", "V1.0，依据当前工程代码与实现结果编写"),
        ("工程根目录", "D:\\Ruanjian\\Vivado\\zynq_chengxu\\ZYNQ_7015\\ZYNQ_7020_yuntu"),
        ("目标平台", "Xilinx Zynq-7020，PL 采集与 FFT，PS 进行波束形成和网络传输"),
        ("适用对象", "系统开发、联调、部署、定位精度优化和故障排查人员"),
        ("当前新增能力", "从 FFT 前提取通道 0 的真实原始时域采样，并在 Qt 上位机同步显示"),
    ], [2100, 7020])

    add_note(doc, "阅读说明", "本文以工程中已经存在的源代码、Block Design 和构建产物为准。未在代码中明确给出的阵列物理尺寸、麦克风灵敏度、dist.bin 的生成坐标系，不作未经验证的假设。", "warning")

    add_heading(doc, "摘要", 1)
    add_text(doc, "本系统用于局部放电声学信号的阵列采集、频域波束形成和可视化定位。PL 端完成多通道数据整理、帧缓存、FFT 供数以及新增的原始时域旁路采集；PS 端接收 FFT 结果，选择有效频点，依据离线生成的距离查找表进行 41 x 41 网格扫描，并将降采样后的 32 x 24 强度云图通过 UDP 发送给 Qt 上位机。上位机叠加云图与相机画面，同时接收并绘制通道 0 的原始时域波形。")
    add_text(doc, "定位无法输出的常见根因并不只在于“峰值作为定位准则”。旧逻辑如果只依赖单通道、固定阈值或每通道独立归一化，容易让宽带、低信噪比或通道幅相不一致的局部放电脉冲失去可用于波束形成的空间信息。当前实现已经改为全阵列平均功率选频、最多叠加 6 个自适应候选频点、宽带回退至全局最强频点，并通过全阵列统一归一化保留通道间相对幅相。")

    add_heading(doc, "目录", 1)
    toc = [
        "1. 系统范围与功能边界", "2. 关键参数与数据规模", "3. 总体架构与数据流", "4. FPGA 端设计", "5. PS 端软件与 DMA 管理",
        "6. 定位与波束形成算法", "7. UDP 通信协议", "8. Qt 上位机设计", "9. 原始时域波形扩展", "10. 构建、导出与下载流程",
        "11. 使用流程与验收标准", "12. 故障定位与排障", "13. 已知限制与优化建议", "附录 A. 工程文件索引", "附录 B. 版本与变更记录"
    ]
    for item in toc:
        add_bullet(doc, item)

    page_break(doc)
    add_heading(doc, "1. 系统范围与功能边界", 1)
    add_heading(doc, "1.1 设计目标", 2)
    for item in [
        "对 64 路声学阵列同步采样数据进行帧级频域处理，形成用于空间定位的复数观测向量。",
        "在指定频段内自动识别主导频点，支持单源和双源定位模式。",
        "对预设扫描区域执行延时求和波束形成，输出云图而不是仅输出一个坐标点。",
        "通过以太网 UDP 将云图、主频信息和原始时域波形送至上位机。",
        "提供从数据链路、频段设置、DMA、网络到上位机显示的完整联调路径。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "1.2 系统不直接保证的内容", 2)
    add_text(doc, "系统输出的热点表示扫描网格中的相对波束形成能量，不等同于经过绝对声压标定后的放电量，也不自动完成摄像机内外参标定、阵列几何误差修正或三维坐标反演。定位精度最终受阵列物理坐标、通道排序、声速、采样同步、声学反射和 dist.bin 的正确性共同决定。")
    add_note(doc, "特别说明", "当前原始时域显示是 FFT 之前的真实采样值，不是根据频谱做反变换得到的波形；但它当前只输出硬件通道 0，长度为 128 点。", "note")

    add_heading(doc, "2. 关键参数与数据规模", 1)
    add_table(doc, ["项目", "当前值", "含义与影响"], [
        ("阵列通道数", "64", "定位使用全部 64 通道的复数 FFT 观测值。"),
        ("每帧采样点", "128", "FFT 长度和时域波形长度均为 128 点。"),
        ("采样率", "192000 Hz", "奈奎斯特频率 96 kHz。"),
        ("频率分辨率", "1500 Hz", "192000 / 128；频点 b 对应 b x 1500 Hz。"),
        ("可用频点", "1 至 63", "跳过直流；当前候选峰搜索进一步限制为 2 至 62。"),
        ("默认频段", "1 kHz 至 96 kHz", "上位机可经 UDP 调整。"),
        ("扫描网格", "41 x 41 = 1681", "PS 端 beamforming 的原始空间网格。"),
        ("传输云图", "32 x 24 = 768 字节", "由 41 x 41 网格抽样缩小后发送。"),
        ("声速", "343.0 m/s", "参与相位补偿；应随温湿度和现场工况校正。"),
        ("单源候选频点", "最多 6", "降低局放宽带能量只落在单一频点的风险。"),
        ("多源候选数", "最多 2", "第二候选功率低于主候选 15% 时不渲染。"),
        ("原始波形", "通道 0，128 x int16", "单帧时长约 0.667 ms，UDP 负载 268 字节。"),
    ], [1920, 1960, 5240])
    add_text(doc, "FFT 主数据的 DMA 接收缓冲区大小为 64 x 128 x 8 = 65536 字节。其每个元素为 64 位字，软件按低 24 位实部、高 24 位虚部解析并进行符号扩展。")

    add_heading(doc, "3. 总体架构与数据流", 1)
    add_heading(doc, "3.1 处理分层", 2)
    add_flow(doc, [
        ("声学前端", "ADAU7118 等多通道音频采集前端提供同步串行采样数据。"),
        ("PL 数据整理", "axis_parser 将连续输入拼装为每拍 64 通道 x 16 位的 frame_data。"),
        ("PL 频域支路", "data_buffer_manager / BRAM / fft_feeder 向 XFFT 提供 128 点帧，FFT 结果进入原有 S2MM DMA。"),
        ("PL 时域旁路", "raw_waveform_capture 从 FFT 前的 frame_data 取通道 0，打包为 64 位 AXI Stream，送入新增 axi_dma_raw_0。"),
        ("PS 运算", "ARM 接收 FFT 帧，完成选频、通道重排、统一归一化、延时求和、能量图生成与 UDP 封包。"),
        ("PC 显示", "Qt 同时解析 CLOUD 和 WAVE 报文，分别绘制相机叠加云图与底部时域波形。"),
    ])
    add_heading(doc, "3.2 两条数据路径", 2)
    add_table(doc, ["路径", "目的", "数据形态", "是否影响定位"], [
        ("频域定位路径", "生成定位云图", "64 通道 x 128 FFT 复数结果", "是，主路径"),
        ("原始时域旁路", "观察采集到的真实 ADC 时域波形", "通道 0 x 128 signed int16", "否，独立 DMA 支路"),
    ], [1800, 2100, 3300, 2160])
    add_note(doc, "隔离原则", "时域旁路不复用原定位 DMA、不改变 FFT 输入，也不改变云图 UDP 结构。这样波形功能异常时，优先检查新增支路，不应直接破坏原有定位路径。", "note")

    page_break(doc)
    add_heading(doc, "4. FPGA 端设计", 1)
    add_heading(doc, "4.1 axis_parser：多通道帧组装", 2)
    add_text(doc, "axis_parser 的输入是 AXI Stream 形式的连续数据。模块内部以 8 个连续传输节拍拼装一个 1024 位 frame_data，随后拉高一个周期的 frame_valid。1024 位数据对应 64 个 16 位采样槽，为下游缓存、FFT 和时域旁路提供完全相同的帧边界。")
    add_code(doc, "frame_data[CHANNEL_INDEX * 16 +: 16]\n// CHANNEL_INDEX = 0 时提取 frame_data[15:0]")
    add_text(doc, "因此，原始时域支路取数点位于 FFT 之前。它看到的是 parser 输出的采样值，而不是 FFT 实部、虚部或能量值。若现场波形完全为零或饱和，首先应检查采集输入和 axis_parser 的 frame_data，而不是定位算法。")
    add_heading(doc, "4.2 原有频域链路", 2)
    add_table(doc, ["模块", "职责", "接口要点"], [
        ("data_buffer_manager", "缓存帧数据并协调下游读取", "接收 frame_valid / frame_data。"),
        ("BRAM", "保存帧级采样数据", "为 FFT 供数提供时序解耦。"),
        ("fft_feeder", "以 FFT IP 所需格式和时序喂入采样", "完成输入配置及帧发送。"),
        ("XFFT", "对各通道 128 点数据变换", "输出复数频谱，软件按 24 位实部和 24 位虚部解释。"),
        ("axi_dma_0", "将 FFT 输出经 S2MM 搬运至 DDR", "中断驱动主定位帧完成。"),
    ], [1900, 3240, 4220])
    add_heading(doc, "4.3 raw_waveform_capture：原始时域采样封装", 2)
    add_text(doc, "新增模块 raw_waveform_capture.v 以 frame_valid 为采样使能。每次 frame_valid 到来时，将指定通道的 16 位值写入 sample_mem[capture_index]；收集 128 个采样后，以 32 个 64 位 AXI Stream beat 输出。每拍按样本 0、1、2、3 到低位至高位的顺序打包，并在第 32 拍同时给出 TLAST。")
    add_table(doc, ["信号", "方向", "说明"], [
        ("clk / rst_n", "输入", "使用现有 PL 处理时钟与低有效复位。"),
        ("frame_valid", "输入", "axis_parser 每次得到 64 通道瞬时帧时有效。"),
        ("frame_data[1023:0]", "输入", "64 个 16 位原始通道采样。"),
        ("m_axis_wave_tdata[63:0]", "输出", "4 个 signed int16 样本打包。"),
        ("m_axis_wave_tvalid / tready", "握手", "DMA 准备好时完成每拍发送。"),
        ("m_axis_wave_tlast", "输出", "stream_index = 31 时表示 128 点帧结束。"),
    ], [2350, 1450, 5560])
    add_note(doc, "工程实现特性", "当 DMA 未准备好时，模块保留最近采集完成的一帧，不会等待导致 FFT 主链路回压。代价是时域显示为抽样帧而不是无损连续录波。", "warning")
    add_heading(doc, "4.4 Block Design 与时钟复位", 2)
    add_text(doc, "连接脚本 scripts/enable_raw_waveform_capture.tcl 在 System.bd 中创建 raw_waveform_capture_0 和 axi_dma_raw_0。新 DMA 仅启用 S2MM、禁用 MM2S 和 Scatter-Gather，数据宽度配置为 64 位。其数据主口接入 axi_smc/S02_AXI，再通过 PS HP0 写入 DDR；AXI-Lite 控制口经 ps7_0_axi_periph/M01_AXI 由 PS GP 时钟域访问。")
    add_text(doc, "数据通路时钟使用 clk_wiz_0/clk_out1；AXI-Lite 使用 processing_system7_0/FCLK_CLK0。为避免复位跨时钟域错误，脚本增加 raw_waveform_lite_reset_0，并将 axi_dma_raw_0/axi_resetn 和 ps7_0_axi_periph/M01_ARESETN 复位到该时钟域。脚本最后执行 validate_bd_design 并保存 BD。")
    add_heading(doc, "4.5 新增 FPGA 文件", 2)
    add_table(doc, ["文件", "用途"], [
        ("AXI_DMA.srcs\\sources_1\\new\\raw_waveform_capture.v", "原始时域采集和 AXI Stream 打包 RTL。"),
        ("scripts\\enable_raw_waveform_capture.tcl", "将 RTL、S2MM DMA、时钟、复位和地址映射加入 Block Design。"),
        ("scripts\\build_raw_waveform_bitstream.tcl", "重新综合、实现、写 bitstream，并导出 System_wrapper.hdf。"),
        ("AXI_DMA.runs\\impl_1\\System_wrapper.bit", "已生成的 PL 配置比特流构建产物。"),
    ], [4380, 4980])

    add_heading(doc, "5. PS 端软件与 DMA 管理", 1)
    add_heading(doc, "5.1 主循环状态机", 2)
    add_text(doc, "main.c 在启动时初始化中断控制器、主 AXI DMA、可选的原始波形 DMA、正弦查找表、SD 卡和 lwIP。随后等待 PC 通过 UDP 发送 START 命令。运行期间主循环持续调用 xemacif_input 处理网络；当 is_locating 为真时，为主 FFT DMA 提交一次 S2MM 接收并等待 RxDone 或 Error。FFT 帧到达后才进入定位处理，并在该处理完成后轮询原始波形 DMA。")
    add_table(doc, ["缓冲区", "类型与大小", "缓存维护", "用途"], [
        ("RX_Buf", "uint64_t[64 x 128]，65536 B", "提交 DMA 前和完成后失效缓存", "保存 FFT 复数结果。"),
        ("Waveform_Buf", "int16_t[128]，256 B", "提交/完成原始 DMA 时失效缓存", "保存通道 0 原始波形。"),
        ("dist_table", "float[1681 x 64]", "从 SD 卡 dist.bin 读入", "每个扫描点到每个通道的距离。"),
        ("energy_map", "float[1681]", "CPU 内存", "保存 41 x 41 扫描能量。"),
    ], [1800, 2580, 2200, 2780])
    add_heading(doc, "5.2 条件编译与 BSP 依赖", 2)
    add_text(doc, "原始波形 DMA 使用 #ifdef XPAR_AXI_DMA_RAW_0_DEVICE_ID 包裹。这样旧 BSP 中没有新 DMA 宏时，原有程序仍能编译并继续输出云图，但不发送 WAVE 报文。只有重新导入更新后的 HDF/XSA、重新生成 BSP，xparameters.h 出现 XPAR_AXI_DMA_RAW_0_DEVICE_ID 后，原始波形功能才会被编译进板端程序。")
    add_note(doc, "部署关键点", "仅替换 Qt 可执行文件或仅替换 C 源码都不会让波形出现；至少需要新 bitstream + 更新 BSP 后重新编译的 ELF + 新 Qt 上位机三者配套。", "warning")
    add_heading(doc, "5.3 原始波形 DMA 服务过程", 2)
    add_code(doc, "if (!raw_wave_dma_active) {\n    XAxiDma_SimpleTransfer(&RawWaveDma, Waveform_Buf, 256, DEVICE_TO_DMA);\n    raw_wave_dma_active = 1;\n} else if (!XAxiDma_Busy(&RawWaveDma, DEVICE_TO_DMA)) {\n    invalidate_cache(Waveform_Buf, 256);\n    send_udp_waveform();\n    raw_wave_dma_active = 0;\n}")
    add_text(doc, "这里采用轮询方式而不是新增中断，规模小、便于验证。由于 raw_waveform_capture 只在主定位帧完成后的服务函数中发起接收，时域帧速率受主定位循环、FFT DMA 完成时间和 CPU 波束形成计算时间限制。它用于可视化与诊断，不应用作高吞吐连续录音链路。")

    page_break(doc)
    add_heading(doc, "6. 定位与波束形成算法", 1)
    add_heading(doc, "6.1 输入频谱格式与通道映射", 2)
    add_text(doc, "软件从 RX_Buf[hw_ch x 128 + bin] 读取每个硬件通道在频点 bin 的 64 位 FFT 数据。低 24 位和 bit[47:24] 分别作为有符号实部、虚部，并手工完成 24 位到 32 位的符号扩展。硬件通道序号随后经 line_idx = hw_ch mod 8、slot_idx = hw_ch / 8、real_ch = line_idx x 8 + slot_idx 重排，以匹配 dist.bin 使用的阵列几何通道顺序。")
    add_heading(doc, "6.2 频点选择：从单通道峰值到全阵列自适应", 2)
    add_text(doc, "当前实现的每个频点功率是 64 个通道复数幅值平方的平均值：P(b) = (1/64) x sum(|X_ch(b)|^2)。这比只检查通道 0 的峰值更稳健，因为局部放电的声波到达不同阵元后幅度会变化，单通道可能刚好处于节点、方向性衰减或噪声污染位置。")
    add_code(doc, "bin_hz = 192000 / 128 = 1500 Hz\nP(b) = (1 / 64) * sum_ch(re_ch(b)^2 + im_ch(b)^2)\nadaptive_threshold = mean(P(b_start ... b_end)) * 1.15")
    add_text(doc, "find_top_n_bins 在用户设置频段内计算功率，选择同时满足局部极大值和自适应阈值的候选。若宽带局放频谱没有严格的局部极大值，则强制保留频段内全局最强频点。因此，“没有峰值”不会直接导致定位完全不输出。单源模式最多选 6 个候选并累积能量；多源模式最多选 2 个候选，第二候选低于主候选功率 15% 时忽略。")
    add_heading(doc, "6.3 保留通道间幅相：全阵列统一归一化", 2)
    add_text(doc, "extract_obs_vector 先累加全部通道功率，再使用 scale = sqrt(64 / total_pwr) 对整个阵列施加同一个缩放因子。这样既避免不同帧总体幅值变化过大，又保留了各通道之间相对幅值和相位关系。若对每个通道单独归一化，会抹掉阵列方向图的重要幅度差异，弱通道噪声也可能被放大，造成波束形成热点漂移或无稳定热点。")
    add_heading(doc, "6.4 延时求和波束形成", 2)
    add_text(doc, "对于扫描网格点 g 和频率 f，波数 k = 2pi f / c。dist_table[g][ch] 给出该网格点到阵元 ch 的传播距离，权值使用 exp(j x k x dist)。软件把观测复数向量与相位补偿相加，得到复合响应 S(g)，再使用 |S(g)|^4 作为能量图强度。多个候选频点的 |S(g)|^4 在同一能量图中累加。")
    add_code(doc, "S(g, f) = sum_ch X_ch(f) * exp(j * k * d(g, ch))\nE(g) = E(g) + |S(g, f)|^4\nk = 2 * pi * f / 343.0")
    add_text(doc, "使用 sine_lut[4096] 和角度索引近似 sin/cos，降低 ARM 上的大量三角函数运算开销。该近似提高了实时性，但距离表、采样率、通道顺序和声速的系统误差仍会比查表量化误差更显著。")
    add_heading(doc, "6.5 定位不到的系统化原因", 2)
    add_table(doc, ["现象", "可能根因", "优先检查"], [
        ("完全没有云图包", "未 START、主 DMA 未完成、网络 IP/端口错误", "串口状态、UDP 8080、RxDone、PC 网卡地址。"),
        ("云图有包但热点随机", "通道顺序与 dist.bin 不匹配、阵列坐标/声速不准、相位不同步", "real_ch 映射、dist.bin 生成脚本、阵列连线和校准。"),
        ("始终无有效候选频点", "频段设置错误、信号过弱、FFT 数据异常", "频段 1 至 96 kHz、原始波形、FFT DMA 数据。"),
        ("热点总在边缘或固定点", "距离表坐标系错误或云图显示映射偏移", "dist.bin、41 x 41 到 32 x 24 映射、相机标定。"),
        ("有声信号但不稳定", "反射、多径、宽带脉冲、环境噪声", "缩窄频段、增加时间积累、提高帧长度或采用宽带方法。"),
    ], [2050, 3420, 3890])

    add_heading(doc, "7. UDP 通信协议", 1)
    add_heading(doc, "7.1 网络地址与端口", 2)
    add_table(doc, ["端点", "IP / 端口", "方向", "用途"], [
        ("Zynq 板端", "192.168.1.10", "双向", "静态 IP。"),
        ("上位机", "192.168.1.12", "接收", "板端云图和波形的目标地址。"),
        ("数据端口", "UDP 8080", "板端 -> PC", "CLOUD 与 WAVE 报文共用端口。"),
        ("控制端口", "UDP 8081", "PC -> 板端", "模式、开始、停止、频段命令。"),
    ], [1620, 2200, 1860, 3680])
    add_heading(doc, "7.2 云图数据报 CLOUD", 2)
    add_text(doc, "CloudPacket_t 使用 packed 布局，总长度固定为 780 字节。发送端采用目标 PC 的 8080 端口；Qt 端仅在长度、magic 均正确时解析。浮点和整数均按当前 Zynq 与 Qt 运行环境的小端序直接 memcpy 传输。跨平台实现时必须显式规定小端字节序。")
    add_table(doc, ["偏移", "字段", "类型", "长度", "说明"], [
        ("0", "magic", "uint32", "4", "0x5A5AA5A5，标识 CLOUD。"),
        ("4", "peak_freq", "float", "4", "主候选频率，单位 Hz。"),
        ("8", "peak_pwr", "float", "4", "能量图最大值，不是声压 dB。"),
        ("12", "grid", "uint8[768]", "768", "32 x 24 相对能量，0 至 255。"),
    ], [1100, 1700, 1500, 1000, 4060])
    add_text(doc, "云图打包时，PS 对 41 x 41 energy_map 做索引抽样：src_x = x x 40 / 31，src_y = y x 40 / 23，然后以该帧最大能量归一化到 0 至 255。Qt 端为匹配屏幕坐标可选择翻转 Y 轴。当前每帧独立归一化，适合观察热点位置，不适合直接比较不同帧的绝对强度。")
    add_heading(doc, "7.3 原始波形数据报 WAVE", 2)
    add_table(doc, ["偏移", "字段", "类型", "长度", "说明"], [
        ("0", "magic", "uint32", "4", "0x57415645，ASCII 为 WAVE。"),
        ("4", "sampling_rate", "float", "4", "当前为 192000.0 Hz。"),
        ("8", "channel", "uint16", "2", "当前固定为 0。"),
        ("10", "sample_count", "uint16", "2", "当前固定为 128。"),
        ("12", "samples", "int16[128]", "256", "原始有符号时域采样，非 FFT 数据。"),
    ], [1100, 1700, 1500, 1000, 4060])
    add_text(doc, "WAVE 报文总长为 268 字节。Qt 端先验证长度和 magic，再读取采样率、通道号和 128 个小端 signed int16。上位机根据采样率计算横轴时长：128 / 192000 = 0.0006667 s，即约 0.667 ms。")
    add_heading(doc, "7.4 控制命令", 2)
    add_table(doc, ["命令", "长度", "字节布局", "板端动作"], [
        ("单源", "4", "uint32 cmd = 0", "current_mode = MODE_SINGLE_SOURCE。"),
        ("多源", "4", "uint32 cmd = 1", "current_mode = MODE_MULTI_SOURCE。"),
        ("开始", "4", "uint32 cmd = 2", "is_locating = 1，开始 DMA/定位/发送。"),
        ("停止", "4", "uint32 cmd = 3", "is_locating = 0。"),
        ("设置频段", "12", "uint32 4 + float min_kHz + float max_kHz", "范围限制 1.0 至 96.0 kHz，且 min < max。"),
    ], [1100, 950, 3550, 3760])

    page_break(doc)
    add_heading(doc, "8. Qt 上位机设计", 1)
    add_heading(doc, "8.1 界面组成", 2)
    add_table(doc, ["组件", "实现", "作用"], [
        ("主视频区", "QCameraViewfinder", "显示 USB/本地相机画面。"),
        ("云图叠加层", "AcousticOverlayWidget", "绘制 32 x 24 热力图和频段、强度等 HUD。"),
        ("控制侧边栏", "QPushButton / QDoubleSpinBox", "开始、停止、频段、场景、偏移等操作。"),
        ("波形面板", "QDockWidget + TimeWaveformWidget", "可停靠的原始时域波形显示。"),
        ("网络接收", "QUdpSocket + 30 ms QTimer", "在 UDP 8080 上处理 CLOUD 和 WAVE。"),
    ], [2100, 2860, 4400])
    add_heading(doc, "8.2 报文解析策略", 2)
    add_text(doc, "updateNetwork 在一个定时周期内读取全部待接收 UDP 报文，并分别保留最新 CLOUD 与最新 WAVE 包。这样网络临时堆积时，上位机优先显示最新状态而不是按历史包逐个渲染，避免延迟越来越大。parseCloudPacket 和 parseWavePacket 均以固定报文长度和 magic 作为第一层保护。")
    add_text(doc, "波形窗口使用自适应纵轴：遍历 128 个样本求最大绝对值，向上取整到 256 的整数倍，再按 90% 绘图区高度缩放。横轴以采样率换算毫秒，顶部注明当前硬件通道。该显示逻辑不会对数据做滤波、整流或 FFT，因此可用来检查直流偏置、削顶、脉冲和噪声形态。")
    add_heading(doc, "8.3 操作与板端的对应关系", 2)
    add_table(doc, ["界面操作", "发出的 UDP 控制", "预期板端结果"], [
        ("PREVIEW Start", "先发频段命令 4，再发 START 命令 2", "定位循环开始，云图与波形开始有机会输出。"),
        ("STOP", "命令 3", "停止新一轮定位 DMA；显示保持最后状态或变为空。"),
        ("Min/Max(kHz)", "命令 4，min/max float", "更新频点搜索范围。"),
        ("Scene: PD", "设定 1 至 96 kHz", "适合先做宽带局放排查；仍需按现场噪声收窄。"),
        ("Raw Waveform Dock", "无控制命令", "显示 WAVE 数据；其出现依赖板端 FPGA/BSP/ELF 已更新。"),
    ], [1900, 3230, 4230])
    add_note(doc, "界面参数含义", "Threshold、X Offset、Y Offset 当前参与上位机显示状态或视觉叠加，但并不改变 PS 端波束形成门限与 dist.bin。不要把界面阈值误认为 FPGA/PS 的定位算法门限。", "warning")

    add_heading(doc, "9. 原始时域波形扩展", 1)
    add_heading(doc, "9.1 为什么必须改 FPGA 数据通路", 2)
    add_text(doc, "原有 PS DMA 接收的是 XFFT 的复数输出，其物理含义已经是频域系数，无法恢复为严格意义上的采集原始波形。即使做 IFFT，也会受到频谱截断、帧处理、格式解释和相位等因素影响，不能替代 ADC 原始数据。因此要在上位机显示原始时域，必须在 FFT 前从 frame_data 单独分出一条 DMA 数据路径。")
    add_heading(doc, "9.2 实施结果", 2)
    for item in [
        "新增 raw_waveform_capture RTL，采样点位于 axis_parser.frame_data 输出处。",
        "新增 axi_dma_raw_0，单向 S2MM、64 位数据宽度，独立写入 DDR。",
        "板端 main.c 通过条件编译启动和轮询新 DMA，并封装 WAVE UDP 报文。",
        "Qt 主窗新增 Raw Waveform 停靠面板，接收并自绘 int16 时域曲线。",
        "Block Design 已通过 validate_bd_design；实现目录中已生成 System_wrapper.bit 构建产物。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "9.3 当前边界和使用方式", 2)
    add_table(doc, ["项目", "当前行为", "影响"], [
        ("通道", "固定硬件通道 0", "只能判断该通道波形；不能直接对比 64 通道同步性。"),
        ("长度", "128 点", "只能看到 0.667 ms 短时片段。"),
        ("连续性", "最近完整帧的抽样发送", "不是无丢帧连续录波。"),
        ("触发条件", "定位主循环已启动且主 FFT DMA 正常完成后服务", "停止定位、主 DMA 出错时不会保证持续发波形。"),
        ("显示目的", "采集链路和信号质量诊断", "不替代示波器或长时间原始数据记录。"),
    ], [1850, 3600, 3910])
    add_heading(doc, "9.4 波形判读建议", 2)
    for item in [
        "波形接近全零：检查麦克风供电、音频接口、I2S/TDM 时序、axis_parser 通道取位和 DMA 是否实际完成。",
        "波形长期满量程或平顶：检查前端增益、输入过载、16 位截位和符号解释。",
        "波形只有固定正/负偏置：检查 ADC 偏置、采集格式是否带符号、数据对齐。",
        "可见瞬态脉冲但定位差：优先检查频段、阵列几何/通道映射、dist.bin、相机坐标映射和反射环境。",
        "多通道幅相校验需求：下一阶段应将 CHANNEL_INDEX 设计为 AXI-Lite 可配置，或增加通道轮询/多通道帧封包。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10. 构建、导出与下载流程", 1)
    add_heading(doc, "10.1 重新构建 FPGA", 2)
    add_number(doc, "在 Vivado 中打开 AXI_DMA.xpr，确认 raw_waveform_capture.v 已位于 sources_1。")
    add_number(doc, "若 Block Design 尚未加入时域旁路，执行 source scripts/enable_raw_waveform_capture.tcl；脚本会创建 DMA、连接时钟复位与地址，并校验 BD。")
    add_number(doc, "执行 source scripts/build_raw_waveform_bitstream.tcl。该脚本重置 synth_1 和 impl_1，运行至 write_bitstream，并在成功后导出 AXI_DMA.sdk/System_wrapper.hdf。")
    add_number(doc, "确认 AXI_DMA.runs/impl_1/System_wrapper.bit 的时间戳为本次构建，且实现状态为 write_bitstream Complete。")
    add_heading(doc, "10.2 更新 SDK/BSP", 2)
    add_number(doc, "在 SDK 中更新或重新创建 System_wrapper_hw_platform_0，导入新的 AXI_DMA.sdk/System_wrapper.hdf。")
    add_number(doc, "重新生成 AXI_DMA_Test_bsp。检查生成的 xparameters.h 含有 XPAR_AXI_DMA_RAW_0_DEVICE_ID。")
    add_number(doc, "重新编译 AXI_DMA_Test。若宏不存在，先不要期待 WAVE 报文，说明硬件平台或 BSP 仍是旧版本。")
    add_number(doc, "下载新的 bitstream 和应用 ELF；如从 SD/QSPI 启动，还需按启动方式重新打包 BOOT.bin 或更新应用镜像。")
    add_heading(doc, "10.3 更新上位机", 2)
    add_number(doc, "使用 Qt 5.15.2/MinGW 或与当前工程一致的套件重新构建 pc_app_qt。")
    add_number(doc, "使用包含 TimeWaveformWidget、WAVE magic 和 parseWavePacket 的新可执行文件。验证构建版本位于 pc_app_qt/build/waveform_verify/release/AcousticImagerQt.exe。")
    add_number(doc, "将 PC 网卡配置为 192.168.1.12/24，确认板端 192.168.1.10 可达，并允许 UDP 8080/8081 穿过防火墙。")
    add_note(doc, "下载顺序建议", "先下载 bitstream，再下载依赖该硬件定义重新编译的 ELF，最后启动新 Qt 上位机。若先运行新 ELF 而 PL 仍是旧 bitstream，新 DMA 寄存器地址将不可靠。", "warning")

    page_break(doc)
    add_heading(doc, "11. 使用流程与验收标准", 1)
    add_heading(doc, "11.1 首次联调流程", 2)
    for step in [
        "上电后连接串口，确认板端打印 System Ready，并确认 dist.bin 从 SD 卡加载成功。",
        "确认 PC 网口为 192.168.1.12，启动新版本 AcousticImagerQt。",
        "单击 PREVIEW Start。上位机先发送频段，再发 START；板端应打印相应控制状态。",
        "观察相机画面上是否出现云图；同时检查 Raw Waveform 面板是否由 Waiting for raw samples 变为曲线。",
        "先用已知声源或测试脉冲验证通道 0 原始波形非零、非削顶，再观察云图热点是否随声源位置移动。",
        "需要局放宽带搜索时选择 PD 或设为 1 至 96 kHz；干扰较强时收窄到明确有效频段。",
    ]:
        add_number(doc, step)
    add_heading(doc, "11.2 最小验收项", 2)
    add_table(doc, ["编号", "验收项", "通过标准"], [
        ("A1", "硬件启动", "SD 卡 dist.bin 加载成功；主程序进入等待 START 状态。"),
        ("A2", "网络控制", "START/STOP 和频段命令在串口有对应状态输出。"),
        ("A3", "云图通信", "PC UDP 8080 能连续收到长度 780、magic 正确的 CLOUD 报文。"),
        ("A4", "定位显示", "上位机叠加层更新，主频随有效频段变化。"),
        ("A5", "波形通信", "PC UDP 8080 收到长度 268、magic 为 WAVE 的报文。"),
        ("A6", "波形真实性", "对测试输入改变时，波形幅度/形状同步变化；不是恒定伪造曲线。"),
        ("A7", "定位一致性", "固定声源位置重复采样时，能量主峰在允许误差内稳定。"),
    ], [800, 2600, 5960])
    add_heading(doc, "11.3 推荐验证记录", 2)
    add_text(doc, "建议每次现场测试记录：阵列和声源相对位置、距离、温度、采样率、频段、声速设置、上位机截图、原始波形截图、CLOUD/WAVE 报文抓包结果、bitstream 与 ELF 的构建日期。定位问题往往来自固件、距表和阵列安装状态不一致，记录这些信息能显著缩短复现时间。")

    add_heading(doc, "12. 故障定位与排障", 1)
    add_heading(doc, "12.1 完全没有定位结果", 2)
    add_table(doc, ["检查顺序", "检查内容", "判断与处理"], [
        ("1", "上位机是否点了 PREVIEW Start", "未开始时 is_locating 为 0，板端不会提交定位 DMA。"),
        ("2", "板端是否加载 dist.bin", "失败时 main() 直接返回，系统不会进入网络定位。"),
        ("3", "PC IP 与端口", "板端目标必须是 192.168.1.12:8080；控制发往 192.168.1.10:8081。"),
        ("4", "主 FFT DMA 状态", "检查 RxDone、Error、中断连接与 DMA 回写数据。"),
        ("5", "频段设置", "频段必须 min < max，且实际有效频点受 1500 Hz 分辨率约束。"),
        ("6", "原始波形", "若波形也无数据，先查采集/FPGA/新 DMA；若波形正常但云图无，查 FFT/选频/距离表。"),
    ], [1100, 2660, 5600])
    add_heading(doc, "12.2 有 WAVE 但没有 CLOUD", 2)
    add_text(doc, "说明新增时域支路至少能工作到板端 UDP 发送，但不能证明 FFT 主路径、候选频点或 beamforming 正常。检查主 DMA 是否有 RxDone；检查 RX_Buf 中 FFT 数是否非零；将频段先放宽到 1 至 96 kHz；打印或抓取 find_top_n_bins 的候选 bin 和功率。若候选始终为 0，重点检查 FFT 输出格式、时钟和选频范围。")
    add_heading(doc, "12.3 有 CLOUD 但没有 WAVE", 2)
    add_text(doc, "这是最常见的新功能未闭环情形。按顺序检查：(1) 板上下载的是否是包含 axi_dma_raw_0 的新 bitstream；(2) BSP xparameters.h 是否定义 XPAR_AXI_DMA_RAW_0_DEVICE_ID；(3) main.c 是否实际编进 #ifdef 块；(4) raw_wave_dma_active 是否能从 0 变为 1 后回到 0；(5) PC 是否使用支持 WAVE 报文的新版 Qt。若旧 BSP 编译，云图仍正常，但 WAVE 代码完全不会参与执行。")
    add_heading(doc, "12.4 云图存在但定位不对", 2)
    for item in [
        "验证硬件通道到 real_ch 的重排是否与 dist.bin 制作时的通道序号一致。通道顺序错一行/一列会直接导致热点镜像、旋转或散乱。",
        "验证阵列实际平面、阵元坐标单位与 dist.bin 使用的坐标系一致；毫米/米混用会造成严重相位误差。",
        "声速 343 m/s 仅为近似值。环境温度、封装和反射会影响有效声速；高频时相位误差更敏感。",
        "确认所有 64 通道采样同步、极性一致。一个通道反相或延迟都会降低相干叠加增益。",
        "观察原始波形是否有削顶、异常直流偏置或明显单通道异常。",
        "在近距离、低反射环境下使用单一已知声源先做标定，避免把多径当作算法问题。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "12.5 频点选择的调试方法", 2)
    add_text(doc, "建议临时在 main.c 中打印 b_start、b_end、max_bin、mean_pwr、adaptive_threshold 和前若干候选 bin。注意频段输入单位是 kHz，而板端内部转为 Hz；实际 bin 对应 1.5 kHz 栅格。例如设置 20 至 96 kHz 时，首个可用频点约为 bin 14，即 21 kHz。若信号能量集中在频带边界或低于 3 kHz，候选函数当前将 bin 限制在 2 至 62，需要根据实际任务调整。")

    add_heading(doc, "13. 已知限制与优化建议", 1)
    add_heading(doc, "13.1 当前已知限制", 2)
    add_table(doc, ["限制", "现状", "风险"], [
        ("时域通道", "仅通道 0", "无法从上位机直接检查全阵列幅相一致性。"),
        ("波形触发", "依赖定位主循环", "定位停止或卡住时，波形不连续。"),
        ("FFT 长度", "128 点", "频率分辨率仅 1.5 kHz，对窄带干扰分离有限。"),
        ("云图强度", "每帧独立归一化", "不能用显示颜色直接比较不同时间的绝对放电强度。"),
        ("扫描结果", "二维 41 x 41 网格", "未显式提供绝对坐标、置信度和三维深度。"),
        ("网络协议", "裸 UDP、小端 float", "丢包、乱序和跨平台字节序未被协议层处理。"),
    ], [2000, 3220, 3940])
    add_heading(doc, "13.2 建议的下一阶段改进", 2)
    for item in [
        "将 raw_waveform_capture 的 CHANNEL_INDEX 改为 AXI-Lite 可配置寄存器，允许上位机选择通道；更进一步可传输 8/64 通道块用于相位校准。",
        "在 PL 加入环形缓冲和触发逻辑，发生阈值脉冲时导出触发前后更长的时域窗口，替代当前抽样帧。",
        "为 raw DMA 配置中断或使用 ping-pong 缓冲，避免轮询造成波形帧率受 CPU 波束形成拖累。",
        "将 FFT 长度升级到 256/512 点，或采用多帧平均、STFT、GCC-PHAT/宽带 SRP 等方法提升宽带局放定位稳定性。",
        "对每个通道做增益、相位、极性和延时标定，在 dist.bin 生成或观测向量中补偿。",
        "把云图主峰转换为阵列坐标，并结合相机标定将其投影到画面；输出坐标、置信度和历史轨迹。",
        "为 UDP 加入序号、时间戳、版本号和 CRC，便于抓包诊断和兼容后续协议演进。",
    ]:
        add_bullet(doc, item)

    page_break(doc)
    add_heading(doc, "附录 A. 工程文件索引", 1)
    add_table(doc, ["文件", "职责", "关键内容"], [
        ("AXI_DMA.sdk\\AXI_DMA_Test\\src\\main.c", "板端应用", "DMA 接收、选频、波束形成、UDP CLOUD/WAVE、控制命令。"),
        ("AXI_DMA.srcs\\sources_1\\new\\axis_parser.v", "PL 数据解析", "将串行/流式采集整理成 64 x 16 位 frame_data。"),
        ("AXI_DMA.srcs\\sources_1\\new\\data_buffer_manager.v", "PL 帧缓存管理", "原有 FFT 数据缓存路径。"),
        ("AXI_DMA.srcs\\sources_1\\new\\fft_feeder.v", "FFT 喂数", "原有 FFT 输入时序与配置。"),
        ("AXI_DMA.srcs\\sources_1\\new\\raw_waveform_capture.v", "新增 RTL", "FFT 前通道 0 原始时域采集与 64 位 AXIS 打包。"),
        ("AXI_DMA.srcs\\sources_1\\bd\\System\\System.bd", "Block Design", "PS、DMA、SmartConnect、时钟复位和新增波形支路连接。"),
        ("scripts\\enable_raw_waveform_capture.tcl", "BD 自动化", "创建/连接 raw_waveform_capture_0 和 axi_dma_raw_0。"),
        ("scripts\\build_raw_waveform_bitstream.tcl", "构建自动化", "生成 bitstream 并导出 HDF。"),
        ("pc_app_qt\\AcousticMainWindow.h/.cpp", "Qt 上位机", "云图、控制、相机和 Raw Waveform 窗口。"),
        ("AXI_DMA.runs\\impl_1\\System_wrapper.bit", "硬件产物", "PL 配置比特流。"),
        ("AXI_DMA.sdk\\System_wrapper.hdf", "硬件定义", "SDK 平台和 BSP 更新输入。"),
    ], [3560, 1880, 3920])
    add_heading(doc, "附录 B. 版本与变更记录", 1)
    add_table(doc, ["版本", "日期", "内容"], [
        ("V1.0", "2026-07-29", "基于当前工程整理系统架构、定位算法、UDP、Qt 上位机、原始时域旁路、构建部署和排障说明。"),
        ("定位逻辑更新", "当前工程", "从单通道峰值倾向改为全阵列平均功率、自适应候选频点、宽带回退、统一归一化和多频点能量累积。"),
        ("时域链路更新", "当前工程", "新增 raw_waveform_capture、axi_dma_raw_0、WAVE UDP 及 Qt 波形停靠面板。"),
    ], [1550, 1750, 6060])
    add_heading(doc, "附录 C. 现场问题快速判断", 1)
    add_text(doc, "当用户描述“什么都定位不到”时，先不要立即继续调峰值准则。先看三个事实：第一，PC 是否收到 CLOUD；第二，通道 0 是否收到并显示非零 WAVE；第三，主 FFT DMA 的 64 x 128 数据是否有效。三者分别把问题划分为网络/控制层、采集/FPGA 层、FFT/定位算法层。先建立这条最短诊断路径，再针对频段、候选峰、通道映射和距离表做精调，效率最高。")

    doc.core_properties.title = "局部放电声学成像系统技术文档"
    doc.core_properties.subject = "Zynq-7020 64 通道局部放电声学定位系统"
    doc.core_properties.author = "工程技术文档"
    doc.core_properties.comments = "根据当前工程代码与构建状态生成。"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
